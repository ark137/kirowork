#!/usr/bin/env python3
"""
生成给品牌方的《脚本与创意说明》：docs/一棵树的秘密-脚本与创意说明.md 和 .docx
镜头表、字幕、时长全部从 src/timeline.ts 读取（timeline 是唯一数据源），时间线变了重跑即可。
需要：node（已装 typescript）、python-docx
"""
import json, os, re, subprocess
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'docs')
NAME = '一棵树的秘密-脚本与创意说明'
FPS = 30

# ───────────── 读取 timeline ─────────────
JS = r"""
const ts=require('typescript'),fs=require('fs');
const src=fs.readFileSync('src/timeline.ts','utf8');
const js=ts.transpileModule(src,{compilerOptions:{module:'commonjs',target:'es2020'}}).outputText;
const m={exports:{}};
new Function('module','exports','require',js)(m,m.exports,()=>({FPS:30}));
process.stdout.write(JSON.stringify(m.exports.SHOTS));
"""
SHOTS = json.loads(subprocess.run(['node', '-e', JS], cwd=ROOT, capture_output=True, text=True, check=True).stdout)

# ───────────── 幕结构（给品牌方看的叙事分段） ─────────────
ACTS = [
    ('序章', [0, 1, 1.5], '以两个问句开场——“一棵树，要长多少年？一个人，又要走多远？”年轮化作落日，片名落笔。',
     '金色圆点、年轮生长、落日、人的剪影'),
    ('第一幕 · 木的诞生', [2, 3], '四亿年前陆地荒芜，植物学会造“骨头”（木质素），地球有了第一片森林。这是“木线”的起点。',
     '苔藓、细胞、古羊齿林、穿过树冠的光'),
    ('第二幕 · 人与木', [3.1, 3.2, 3.3, 3.4, 3.5], '从火到独木舟、榫卯、宫殿斗拱，再到一盏灯下的木屋：几千年来，人始终与木头为伴——温暖、柔韧、让人安心。落点是江南木屋，顺势进入主人公的故乡温州。',
     '篝火、独木舟、榫卯金线、斗拱堆叠、雨夜木屋暖光'),
    ('第三幕 · 人的起点', [4, 5, 5.5, 6, 7], '1966 出生于温州；学船舶设计；做六年验船师；1994 年辞去公职，28 平方米的小店开张；亲手设计“地板堆成的大树”商标。这是“人线”的起点。',
     '水墨老街、船舶图纸、渔港验船、首店、金色地板条堆成树'),
    ('第四幕 · 断粮', [10, 11, 12], '1998 年国家禁止砍伐，原料来路被斩断；想从巴西进口，货源却攥在中间商手里；他混进码头搬运工里，一箱一箱找线索。',
     '红章公文、被掐住的供应链、夜色码头与手电光'),
    ('第五幕 · 一纸传真', [13, 14, 15], '一卷封箱胶带上的电话号码；语言不通、11 小时时差，他一直拨；一个多月后传真机响了，巴西供应商路易斯寄来报价单。',
     '胶带特写、温州与巴西两个钟面、传真纸、亮起的传真机'),
    ('第六幕 · 雨林里的树', [8, 9], '报价单来自南美——数千万年的雨水、阳光和河流，养出了亚马逊。树为了抢一缕阳光，长成各自的样子。',
     '大陆漂移、雨林俯瞰、云雾、各具形态的热带树'),
    ('第七幕 · 相遇', [16, 17, 17.5, 18], '1999 年，两条时间线在雨林交汇：四亿年长出的树，等来了找它的人。他与路易斯在雨林边建起原材料工厂，“先规划，后采伐”，卫星把林区分成 25 块，每年只采一块、采完就种。',
     '手扶巨树、年轮与时间轴并线、树种名牌、握手、规划图、金色网格、卫星轮伐'),
    ('第八幕 · 走出去', [19, 20, 21], '1999 年在上海青浦建厂；从温州到美洲再到非洲；奥运场馆、上海世博中心、整木定制，成为许多重要地方的脚下。',
     '锯齿厂房、世界地图时间轴、抽象建筑线稿'),
    ('第九幕 · 传承', [22, 23, 24], '2017 年卢奕开留学十年后回到公司；父子常在会上争得面红耳赤，父亲说“只要儿子说得对，就该有心胸接受”；2022 年接任总裁，“从冲突到理解，互为铠甲”。',
     '机场、会议桌两端的剪影、两条缠绕成一棵树的枝干、里程碑'),
    ('尾声', [25, 26, 27], '信任超越任何语言；放大的 Logo 变成一扇扇生活的窗；年轮长回金色大树，收束在“安信地板”和一个“信”字。',
     '书页、Logo 窗格里的家、年轮长成金树'),
]

byid = {s['id']: s for s in SHOTS}
assert sorted(i for _, ids, _, _ in ACTS for i in ids) == sorted(byid), '有镜头未归入幕'


def tc(sec):
    """m:ss；带半秒的时间点写成 m:ss.5"""
    if abs(sec - round(sec)) > 0.01:
        m, r = divmod(sec, 60)
        return f'{int(m)}:{r:04.1f}'
    sec = round(sec)
    return f'{sec // 60}:{sec % 60:02d}'


# 顺序与起点（按 timeline 顺序）
START, cur, SEQ = {}, 0.0, {}
for n, s in enumerate(SHOTS, 1):
    START[s['id']] = cur
    SEQ[s['id']] = n
    cur += s['dur']
TOTAL = cur
assert all(s['id'] in START for s in SHOTS)
act_time = []
for name, ids, intent, kw in ACTS:
    a = min(START[i] for i in ids)
    b = max(START[i] + byid[i]['dur'] for i in ids)
    act_time.append((a, b))

def N(*ids):
    """镜头 id → 第 n 镜（成片中的顺序号）；多个 id 取首尾"""
    a, b = SEQ[ids[0]], SEQ[ids[-1]]
    return f'第 {a} 镜' if a == b else f'第 {a}–{b} 镜'


# 字幕里的绝对化用语自查
flagged = []
for s in SHOTS:
    for sub in s['subs']:
        for m in re.finditer(r'最|第一(?!片)|唯一|绝对', sub['zh']):
            flagged.append((SEQ[s['id']], sub['zh']))

LINE = {'A': '木线（地球时间）', 'B': '人线', 'AB': '两线交汇', 'H': '人与木（人类历史时间）', '—': ''}

# ───────────── 文档内容（块结构，同时渲染成 md 和 docx） ─────────────
B = []  # ('h1'|'h2'|'h3'|'p'|'bullets'|'table'|'quote', payload)


def h1(t): B.append(('h1', t))
def h2(t): B.append(('h2', t))
def h3(t): B.append(('h3', t))
def p(t): B.append(('p', t))
def bl(items): B.append(('bullets', items))
def table(head, rows, widths=None): B.append(('table', (head, rows, widths)))
def quote(t): B.append(('quote', t))


B.append(('title', '《一棵树的秘密》'))
B.append(('subtitle', '安信地板品牌动画 · 脚本与创意说明'))
p(f'版本日期：2026-10-07　｜　全片时长：{tc(TOTAL)}（{TOTAL:g} 秒）　｜　{len(SHOTS)} 个镜头、11 个段落　｜　1920×1080 · 30fps · 中英双语字幕')

h1('一、项目概览')
table(['项目', '说明'], [
    ['片名', '《一棵树的秘密》　副标题：安信地板的前世今生'],
    ['一句话', '树的秘密，是时间；人的秘密，是一个“信”字。'],
    ['时长 / 规格', f'{tc(TOTAL)}（{TOTAL:g} 秒），1920×1080，30 帧/秒，中英双语字幕'],
    ['形式', '2D 水墨风格矢量动画（Remotion + SVG 全程制作），无真人实拍、无真实人物肖像'],
    ['声音', '全片拟音与环境声（无音乐、无旁白），每个声音与画面动作一一对齐'],
    ['内容范围', '从四亿年前第一片森林，到创始人卢伟光的人生，到安信地板的全球布局与父子传承'],
], [3, 13])

h1('二、创意主旨')
quote('一棵树，要长多少年？一个人，又要走多远？')
p('全片用一个比喻贯穿：**安信是一棵树**。树要慢慢长，信赖也要慢慢长。片名里的“秘密”有两层答案，都在片尾揭晓：树的秘密是“时间”，人的秘密是“信”。')
h3('两条时间线，在雨林里相遇')
bl([
    '**木线（地球时间）**：四亿年前的木质素 → 第一片森林 → 亚马逊雨林。墨色，带年轮小图标，前缀 EARTH TIME。',
    '**人线**：卢伟光 1966 → 1994 → 1999 → 今天。金色，是全片的情感主线。',
    '**人与木（人类历史时间）**：介于两者之间的一段——人类几千年与木头共处的历史。赭石色，前缀 HUMAN TIME。',
    '**1999 年，两条线在亚马逊雨林交汇**：巨树的年轮与一根金色时间轴并到一起（“相遇”一镜），之后全片进入“人把木头带进千家万户”的下半程，最后收束到 Logo。',
])
h3('贯穿全片的三个视觉母题')
bl([
    '**年轮**：时间的可视化。开场年轮化作落日，“相遇”里年轮与时间轴并线，结尾年轮长回金色大树。',
    '**堆叠**：安信商标是“地板堆成的大树”。片中反复出现同一个动作——一块一块叠起来：地板条叠成 Logo（' + N(7) + '）、故宫斗拱一层层叠起（第二幕）、Logo 变成一扇扇窗（尾声）。“信赖要像树一样慢慢长成”，正是靠一层层叠出来的。',
    '**金色**：企业 VI 主色。只在“人”与“信”的关键时刻出现——商标里的金色地板条、榫卯合上时的金线、相遇时掌心的金色涟漪、规划线变成金色网格、最后的金色大树。',
])
h3('叙事节奏')
p(f'全片 {tc(TOTAL)} 分成 11 个段落。前 {tc(act_time[3][1])} 建立“木”与“人”的世界，落在亲手设计的商标上；中段（{tc(act_time[4][0])}–{tc(act_time[7][1])}）是全片的戏剧高潮——断粮、一纸传真、雨林相遇、建厂与轮伐；后段（{tc(act_time[8][0])}–{tc(TOTAL)}）讲扩张与传承，用“信”收束。')

h1('三、视觉与声音')
h3('视觉语言')
bl([
    '**宣纸水墨**：宣纸底色与纤维肌理，墨色洇染边缘，淡墨远景与浓墨前景形成层次；夜景也保持中等明度，不压黑。',
    '**色板**：宣纸米、墨、企业金、朱红（印章与“安信”红字）、自然绿。',
    '**人物**：全部是剪影，不做肖像。卢伟光为及肩长发，路易斯戴草帽，卢奕开为短发；风格统一，不涉及真人面貌。',
    '**Logo**：按原图逐块识别重建（46 块地板条），可逐块动画。首家门店按原照片用插画重绘。',
    '**镜头**：以缓推、缓拉和横移为主；场景之间 14 帧（约半秒）交叉淡化；每一镜只做一个视觉动作，画面干净。',
    '**字幕**：底部双语，中文逐字墨晕显影，英文为斜体衬线；暗底镜头自动换浅色。',
    '**年份标签**：左上角，三种样式对应三种时间（见上）。',
])
h3('声音设计')
bl([
    '**不加音乐、不加旁白**，由拟音与环境声带节奏：风、雨林虫鸟、水声、柴火、木头敲击、榫卯合拢的“咔”、电话拨号音、传真握手音、钟声与风铃。',
    '素材来自 Kenney 音效包（CC0 协议，可商用），其余为程序合成。',
    '整体响度偏安静，为后续叠加配乐或旁白留出空间。',
])

h1('四、分镜脚本')
p('下表按时间顺序。“时间”为该镜在全片中的起点；“年份标签”为画面左上角的文字；字幕为中文 / 英文。')

for (name, ids, intent, kw), (a, b) in zip(ACTS, act_time):
    h2(f'{name}　{tc(a)}–{tc(b)}（{b - a:g} 秒）')
    p(f'**本段意图**：{intent}')
    p(f'**画面关键词**：{kw}')
    rows = []
    for i in ids:
        s = byid[i]
        tag = ' · '.join(x for x in [s.get('year', ''), s.get('yearNote', '')] if x) or '—'
        subs = '\n'.join(f'{x["zh"]}\n{x["en"]}' for x in s['subs']) or '（无字幕）'
        rows.append([str(SEQ[i]), tc(START[i]), f'{s["dur"]:g}s', tag, s['brief'], subs])
    table(['#', '时间', '时长', '年份标签', '画面', '字幕（中 / 英）'], rows, [0.9, 1.3, 1.2, 3.2, 8.3, 8.6])

h1('五、全片字幕文案（连读）')
p('方便整体审读文案，不含画面说明。')
for name, ids, _, _ in ACTS:
    B.append(('para_group', (name, [byid[i] for i in ids])))

h1('六、请品牌方确认的事项')
p('以下分四类（A 事实数据、B 史实地理、C 文案合规、D 画面处理）。标注“品牌方提供”的内容，来自此前提供的资料，正式发布前请再核一遍；标注“制作方补充”的内容，是制作时为串起故事加入的，需要您或文史顾问确认。')

h3('A. 事实与数据（品牌方提供）')
table(['镜头', '内容', '需确认'], [
    [N(4, 7), '1966 年出生温州经商世家；大学学船舶设计（标签写 DALIAN UNIV. OF TECHNOLOGY）；1988–1994 任中华人民共和国船舶检验局温州渔船检验处验船师（2026-10-07 已按品牌方最新提供修订）；1994 年辞去公职；1994.4.8 首店 28 ㎡开业', '时间、面积；机构中文全称已按最新提供修改，英文译名 PRC Ship Inspection Bureau\'s Wenzhou Fishing Vessel Inspection Office 为制作方翻译，请确认'],
    [N(10, 15), '1998 年禁伐；想从巴西进口；1998 年 7 月在封箱胶带上找到号码；11 小时时差；一个多月后收到报价单', '时间点、“11 小时”、“一个多月”'],
    [N(17.5) + ' / ' + N(18), '与路易斯共建原材料工厂；林区分 25 块、每年只采一块、采完就种，25 年一个轮回', '是否为实际作业方式；路易斯姓名及形象使用是否已征得同意'],
    [N(19, 21), '1999 年上海青浦工厂；2004–2006 拉美四国、收购美国 ARK FLOORS、凯雷注资 5000 万美金、投资非洲木材；2007–2010 奥运场馆、苏州吴江工厂、世博中心、整木定制', '年份、金额、机构全称与对外披露口径（投融资信息尤其需确认是否可公开）'],
    [N(22, 24), '2017 年卢奕开回公司；2022 年任总裁；2023 年“树木的秘密”展厅；2025 年上海旗舰店焕新', '时间、职务'],
    [N(23), '引用卢伟光的话：“只要儿子说得对，就该有心胸接受，否则，企业成不了百年企业。”', '本人是否同意以这种方式引用'],
    [N(25), '出现多本书刊的书页：《读者》、MBA 教材、国家统编中学生教材、浙江省初中德育课程教材、《风云浙商-2》', '收录情况属实；画面若出现真实书名或封面，需确认引用授权'],
    ['全片', '英文名使用 Carl Lu、Ben Lu', '是否为官方英文名'],
], [3, 10.5, 6.5])

h3('B. 史实与地理（制作方补充）')
table(['镜头', '表述', '说明'], [
    [N(2), '四亿年前陆地荒芜；植物造出木质素', '泥盆纪时期，科学共识'],
    [N(3), '约 3.85 亿年前第一批森林', '以古羊齿（Archaeopteris）为代表的最早树木，常见表述'],
    [N(3.1), '“数十万年前”人类用火', '学界对用火起始年代有不同说法，这里用宽泛表述；如要更严谨可改成“远古”'],
    [N(3.2), '约 8000 年前，浙江先民把原木凿成独木舟（标签：跨湖桥）', '杭州萧山跨湖桥遗址的独木舟，常见年代为约 8000 年前'],
    [N(3.3), '约 7000 年前，河姆渡人发明榫卯', '余姚河姆渡遗址出土的榫卯木构件，常见年代为约 7000 年前'],
    [N(3.4), '1420 年，紫禁城用一层层斗拱托起宫殿屋檐', '紫禁城于明永乐十八年（1420）建成，为世界现存规模最大的木结构宫殿建筑群之一'],
    [N(8), '亚马逊是地球上最大的热带雨林', '地理常识'],
], [3, 8, 9])
p('建议上线前请文史顾问或品牌方的文化部门过一遍上表。')

h3('C. 文案措辞与合规')
bl([
    '**已避免绝对化用语**：我们没有使用“最好的建筑材料”这类说法，改成“温暖、柔韧、让人安心”。片中含“最”字的句子只有：' + '、'.join(f'第 {n} 镜“{z}”' for n, z in flagged) + '。其中前两处是时间或地理的客观描述，第三处见下一条。',
    '**' + N(17) + '“找到了最适合的木头”**：是叙事语气，不是产品宣称，但含“最”字。英文为 the perfect wood，同样偏绝对。保守起见可改为“找到了合适的木头”（英文 the right wood），请品牌方决定。',
    '**“更温暖、更健康的生活空间”（' + N(26) + '）**：这是品牌愿景式表达。如用于对外投放，建议确认“健康”一词是否有检测或认证支撑。',
    '**“只取森林给得起的那一份”“森林生生不息”**：属可持续理念表达。若用于公开发布，建议与实际的认证（如 FSC 等）或数据对应，由品牌方确认，制作方不另行添加任何认证标志。',
])

h3('D. 画面处理')
bl([
    '**片名**：片内标题卡为《一棵树的秘密》，工程早期曾用《一棵树的信》，请确认最终片名。',
    '**“卢记”布招**（' + N(4) + '，1966 温州老街）：是营造“经商世家”氛围的艺术处理，并非史实依据。如不希望暗示具体店号，可以换成通用招牌字。',
    '**建筑线稿**（' + N(21) + '）：奥运场馆、世博中心等为抽象线稿，不是真实建筑外观。',
    '**世界地图**（' + N(20) + '）：画面里的国家标注请品牌方对照实际业务核对一遍。',
    '**人物**：全片为剪影，无面部，不构成对真实人物的肖像再现。',
])

h1('七、制作状态')
bl([
    f'全部 {len(SHOTS)} 个镜头已制作完成，每一镜均有对应的音效。',
    '早期完整成片（30 个镜头，203 秒，含音效）已交付；之后新增的两段——“雨林原材料工厂”（8 秒）与“人与木”（五镜，29 秒）——以单独片段交付，首尾与原片逐帧衔接，可直接插入。',
    f'**合并后的 {TOTAL:g} 秒完整版尚未整体渲染**。文案和分镜确认后，可以一次性出完整成片。',
    '**修订记录（2026-10-07）**：' + N(5.5) + '（验船师）机构名称改为“中华人民共和国船舶检验局温州渔船检验处”，字幕与证书上的机构名同步修改；因字幕加长，该镜由 9 秒延长到 10.5 秒，全片相应增加 1.5 秒。该镜已单独交付替换片段。',
    '目前**没有旁白和音乐**。如需，可在文案定稿后再做：旁白可直接用本文第五节的中文字幕，配乐可按本文的段落分段设计。',
    '如需缩短版本（例如 60 秒或 90 秒精华版），可以从本表中挑选镜头重新编排，请告知使用场景。',
])

# ───────────── 渲染：Markdown ─────────────
md = []
for kind, x in B:
    if kind == 'title':
        md.append(f'# {x}\n')
    elif kind == 'subtitle':
        md.append(f'**{x}**\n')
    elif kind == 'h1':
        md.append(f'\n---\n\n## {x}\n')
    elif kind == 'h2':
        md.append(f'\n### {x}\n')
    elif kind == 'h3':
        md.append(f'\n#### {x}\n')
    elif kind == 'p':
        md.append(x + '\n')
    elif kind == 'quote':
        md.append(f'> {x}\n')
    elif kind == 'bullets':
        md.append('\n'.join(f'- {i}' for i in x) + '\n')
    elif kind == 'table':
        head, rows, _ = x
        md.append('| ' + ' | '.join(head) + ' |')
        md.append('|' + '|'.join(['---'] * len(head)) + '|')
        for r in rows:
            md.append('| ' + ' | '.join(c.replace('\n', '<br>') for c in r) + ' |')
        md.append('')
    elif kind == 'para_group':
        name, shots = x
        txt = [sub['zh'] for s in shots for sub in s['subs']]
        en = [sub['en'] for s in shots for sub in s['subs']]
        md.append(f'**{name}**\n')
        if txt:
            md.append('\n\n'.join(f'{a}  \n*{b}*' for a, b in zip(txt, en)) + '\n')
        else:
            md.append('（片名卡：《一棵树的秘密》——安信地板的前世今生）\n')

os.makedirs(OUT, exist_ok=True)
md_path = os.path.join(OUT, NAME + '.md')
open(md_path, 'w', encoding='utf-8').write('\n'.join(md))

# ───────────── 渲染：Word ─────────────
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
for side in ('left_margin', 'right_margin'):
    setattr(sec, side, Cm(1.8))
sec.top_margin = sec.bottom_margin = Cm(1.6)
FONT = 'Microsoft YaHei'


def set_font(run, size=10.5, bold=False, color=None, italic=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_rich(par, text, size=10.5, color=None):
    """**粗体** 简易标记"""
    for k, part in enumerate(re.split(r'\*\*', text)):
        if part:
            set_font(par.add_run(part), size, bold=(k % 2 == 1), color=color)


def shade(cell, hexcolor):
    tcPr = cell._element.get_or_add_tcPr()
    from docx.oxml import OxmlElement
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


GOLD = (0xB9, 0x79, 0x0F)
INK = (0x2A, 0x24, 0x1E)
for kind, x in B:
    if kind == 'title':
        par = doc.add_paragraph()
        set_font(par.add_run(x), 28, True, INK)
    elif kind == 'subtitle':
        par = doc.add_paragraph()
        set_font(par.add_run(x), 14, True, GOLD)
    elif kind == 'h1':
        par = doc.add_paragraph()
        par.paragraph_format.space_before = Pt(16)
        set_font(par.add_run(x), 17, True, GOLD)
    elif kind == 'h2':
        par = doc.add_paragraph()
        par.paragraph_format.space_before = Pt(12)
        par.paragraph_format.keep_with_next = True
        set_font(par.add_run(x), 13.5, True, INK)
    elif kind == 'h3':
        par = doc.add_paragraph()
        par.paragraph_format.space_before = Pt(8)
        par.paragraph_format.keep_with_next = True
        set_font(par.add_run(x), 12, True, GOLD)
    elif kind == 'p':
        par = doc.add_paragraph()
        add_rich(par, x)
    elif kind == 'quote':
        par = doc.add_paragraph()
        par.paragraph_format.left_indent = Cm(0.8)
        set_font(par.add_run(x), 15, True, GOLD)
    elif kind == 'bullets':
        for item in x:
            par = doc.add_paragraph(style='List Bullet')
            add_rich(par, item)
    elif kind == 'table':
        head, rows, widths = x
        t = doc.add_table(rows=1, cols=len(head))
        t.style = 'Table Grid'
        t.autofit = False
        for j, h in enumerate(head):
            c = t.rows[0].cells[j]
            c.text = ''
            set_font(c.paragraphs[0].add_run(h), 9.5, True, (0xFF, 0xFF, 0xFF))
            shade(c, '8E5F17')
        for r in rows:
            cells = t.add_row().cells
            for j, v in enumerate(r):
                cells[j].text = ''
                lines = v.split('\n')
                for n, ln in enumerate(lines):
                    par = cells[j].paragraphs[0] if n == 0 else cells[j].add_paragraph()
                    is_en = kind == 'table' and head[-1].startswith('字幕') and j == len(r) - 1 and n % 2 == 1
                    set_font(par.add_run(ln), 8.5 if is_en else 9, italic=is_en, color=(0x6B, 0x5C, 0x4A) if is_en else None)
        if widths:
            total = sum(widths)
            for row in t.rows:
                for j, w in enumerate(widths):
                    row.cells[j].width = Cm(25.9 * w / total)
        doc.add_paragraph()
    elif kind == 'para_group':
        name, shots = x
        par = doc.add_paragraph()
        par.paragraph_format.space_before = Pt(6)
        set_font(par.add_run(name), 11, True, GOLD)
        any_sub = False
        for s in shots:
            for sub in s['subs']:
                any_sub = True
                par = doc.add_paragraph()
                par.paragraph_format.space_after = Pt(0)
                set_font(par.add_run(sub['zh']), 10.5)
                par = doc.add_paragraph()
                set_font(par.add_run(sub['en']), 9, color=(0x6B, 0x5C, 0x4A), italic=True)
        if not any_sub:
            par = doc.add_paragraph()
            set_font(par.add_run('（片名卡：《一棵树的秘密》——安信地板的前世今生）'), 10.5)

docx_path = os.path.join(OUT, NAME + '.docx')
doc.save(docx_path)
print(json.dumps({'md': md_path, 'docx': docx_path, 'shots': len(SHOTS), 'total_s': TOTAL, 'flagged': flagged}, ensure_ascii=False))
